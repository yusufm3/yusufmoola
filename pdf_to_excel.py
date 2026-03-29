"""
pdf_to_excel.py
---------------
Iterate over all PDF files in a given directory (quotes & invoices),
extract line-item tables, and write a structured Excel workbook.

Usage:
    python pdf_to_excel.py <pdf_directory> [output_excel_path]

    pdf_directory      – folder that contains the PDF files
    output_excel_path  – (optional) path for the Excel output
                         defaults to <pdf_directory>/output.xlsx

Requirements:
    pip install pdfplumber openpyxl
"""

import os
import re
import sys
import argparse
from pathlib import Path

import pdfplumber
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
ALT_ROW_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
BODY_FONT = Font(name="Calibri", size=10)
BOLD_FONT = Font(name="Calibri", size=10, bold=True)

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

# Column names that strongly indicate a line-item table
LINE_ITEM_KEYWORDS = {
    "description", "desc", "item", "product", "service", "particulars",
    "qty", "quantity", "units", "unit", "amount", "price", "rate", "total",
    "subtotal", "vat", "tax", "discount", "code", "ref", "part",
}


def _looks_like_header_row(row: list) -> bool:
    """Return True if at least two cells in *row* match known line-item keywords."""
    if not row:
        return False
    matches = sum(
        1
        for cell in row
        if cell and any(kw in str(cell).lower() for kw in LINE_ITEM_KEYWORDS)
    )
    return matches >= 2


def _clean_cell(value) -> str:
    """Strip whitespace from a cell value."""
    if value is None:
        return ""
    return str(value).strip()


def _is_mostly_empty(row: list) -> bool:
    return all(_clean_cell(c) == "" for c in row)


def _extract_tables_from_pdf(pdf_path: str) -> list[dict]:
    """
    Open *pdf_path* with pdfplumber and return a list of table records.

    Each record is a dict:
        {
            "page":    int,
            "headers": [str, ...],
            "rows":    [[str, ...], ...],
        }

    Strategy
    --------
    1. Try pdfplumber's built-in table detector on every page.
    2. If no table is found on a page, fall back to raw text lines and
       attempt to identify a tabular section by looking for header keywords.
    """
    results = []
    with pdfplumber.open(pdf_path) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables()
            if tables:
                for raw_table in tables:
                    if not raw_table:
                        continue
                    # Find which row is the header
                    header_idx = 0
                    for i, row in enumerate(raw_table):
                        if _looks_like_header_row(row):
                            header_idx = i
                            break

                    headers = [_clean_cell(c) for c in raw_table[header_idx]]
                    data_rows = []
                    for row in raw_table[header_idx + 1 :]:
                        if _is_mostly_empty(row):
                            continue
                        data_rows.append([_clean_cell(c) for c in row])

                    if data_rows:
                        results.append(
                            {"page": page_num, "headers": headers, "rows": data_rows}
                        )
            else:
                # Fallback: parse raw text lines into a pseudo-table
                text = page.extract_text() or ""
                lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
                table = _parse_text_lines(lines)
                if table:
                    results.append({"page": page_num, **table})

    return results


def _parse_text_lines(lines: list[str]) -> dict | None:
    """
    Heuristically convert plain text lines into a table dict
    ``{"headers": [...], "rows": [[...], ...]}``.

    Looks for a header-like line then captures subsequent lines until a
    totals keyword is seen or the block ends.
    """
    header_idx = None
    for i, line in enumerate(lines):
        parts = re.split(r"\s{2,}|\t", line)
        if _looks_like_header_row(parts):
            header_idx = i
            break

    if header_idx is None:
        return None

    header_line = lines[header_idx]
    headers = [p.strip() for p in re.split(r"\s{2,}|\t", header_line) if p.strip()]

    STOP_WORDS = {"subtotal", "sub total", "total", "grand total", "balance due"}
    rows = []
    for line in lines[header_idx + 1 :]:
        low = line.lower()
        if any(sw in low for sw in STOP_WORDS):
            break
        parts = [p.strip() for p in re.split(r"\s{2,}|\t", line) if p.strip()]
        if not parts:
            continue
        # Pad or truncate to match header length
        while len(parts) < len(headers):
            parts.append("")
        rows.append(parts[: len(headers)])

    if not rows:
        return None

    return {"headers": headers, "rows": rows}


# ---------------------------------------------------------------------------
# Excel writing
# ---------------------------------------------------------------------------

def _style_header_row(ws, row_num: int, num_cols: int):
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def _style_data_row(ws, row_num: int, num_cols: int, alternate: bool):
    fill = ALT_ROW_FILL if alternate else PatternFill()
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.fill = fill
        cell.font = BODY_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def _auto_fit_columns(ws):
    for col_cells in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            try:
                length = len(str(cell.value)) if cell.value else 0
                if length > max_length:
                    max_length = length
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_length + 4, 50)


def _safe_sheet_name(name: str, used: set) -> str:
    """Return a worksheet name that is ≤31 chars and unique within *used*."""
    name = re.sub(r'[\\/*?:\[\]]', "_", name)[:31]
    base = name
    counter = 1
    while name in used:
        suffix = f"_{counter}"
        name = base[: 31 - len(suffix)] + suffix
        counter += 1
    used.add(name)
    return name


def build_excel(pdf_dir: str, output_path: str):
    pdf_dir = Path(pdf_dir).resolve()
    if not pdf_dir.is_dir():
        print(f"ERROR: '{pdf_dir}' is not a valid directory.")
        sys.exit(1)

    pdf_files = sorted(pdf_dir.rglob("*.pdf"))
    if not pdf_files:
        print(f"No PDF files found in '{pdf_dir}'.")
        sys.exit(0)

    wb = openpyxl.Workbook()
    # Remove default sheet; we'll create our own
    wb.remove(wb.active)

    # Summary sheet (written last once we know all data)
    summary_data = []  # list of (filename, doc_type, num_tables, num_rows)

    used_names: set = set()

    for pdf_path in pdf_files:
        rel_name = pdf_path.relative_to(pdf_dir)
        stem = pdf_path.stem
        print(f"Processing: {rel_name}")

        tables = _extract_tables_from_pdf(str(pdf_path))

        # Guess document type from filename (use word-boundary patterns to avoid
        # false matches like "inventory", "misquote", etc.)
        fname_lower = pdf_path.name.lower()
        if re.search(r'\binvoice\b|\binv\b', fname_lower):
            doc_type = "Invoice"
        elif re.search(r'\bquote\b|\bquotation\b|\brfq\b', fname_lower):
            doc_type = "Quote"
        else:
            doc_type = "Document"

        total_rows = sum(len(t["rows"]) for t in tables)
        summary_data.append((str(rel_name), doc_type, len(tables), total_rows))

        if not tables:
            print(f"  ⚠  No line-item tables found.")
            sheet_name = _safe_sheet_name(stem, used_names)
            ws = wb.create_sheet(title=sheet_name)
            ws.append([f"No extractable tables found in: {rel_name}"])
            ws["A1"].font = BOLD_FONT
            continue

        sheet_name = _safe_sheet_name(stem, used_names)
        ws = wb.create_sheet(title=sheet_name)

        current_row = 1

        # File header banner
        ws.cell(row=current_row, column=1).value = f"{doc_type}: {rel_name}"
        ws.cell(row=current_row, column=1).font = Font(
            name="Calibri", size=13, bold=True, color="1F4E79"
        )
        ws.merge_cells(
            start_row=current_row,
            start_column=1,
            end_row=current_row,
            end_column=max(len(t["headers"]) for t in tables) or 1,
        )
        current_row += 2

        for t_idx, table in enumerate(tables, start=1):
            headers = table["headers"]
            rows = table["rows"]
            num_cols = max(len(headers), max((len(r) for r in rows), default=0))

            # Normalise column count
            while len(headers) < num_cols:
                headers.append("")
            rows = [r + [""] * (num_cols - len(r)) for r in rows]

            # Table label
            label = f"Page {table['page']} – Table {t_idx}"
            ws.cell(row=current_row, column=1).value = label
            ws.cell(row=current_row, column=1).font = Font(
                name="Calibri", size=10, italic=True, color="595959"
            )
            current_row += 1

            # Header row
            for col_idx, header in enumerate(headers, start=1):
                ws.cell(row=current_row, column=col_idx).value = header
            _style_header_row(ws, current_row, num_cols)
            ws.row_dimensions[current_row].height = 20
            current_row += 1

            # Data rows
            for row_idx, row in enumerate(rows):
                for col_idx, value in enumerate(row, start=1):
                    ws.cell(row=current_row, column=col_idx).value = value
                _style_data_row(ws, current_row, num_cols, alternate=(row_idx % 2 == 0))
                current_row += 1

            current_row += 1  # blank row between tables

        _auto_fit_columns(ws)
        ws.freeze_panes = "A3"

    # ----- Summary sheet -----
    ws_sum = wb.create_sheet(title="Summary", index=0)
    ws_sum.append(["PDF Document Extraction Summary"])
    ws_sum["A1"].font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    ws_sum.merge_cells("A1:D1")
    ws_sum.append([])
    ws_sum.append(["File", "Type", "Tables Found", "Total Line Items"])
    _style_header_row(ws_sum, 3, 4)
    ws_sum.row_dimensions[3].height = 20

    for i, (fname, dtype, ntables, nrows) in enumerate(summary_data, start=4):
        ws_sum.cell(row=i, column=1).value = fname
        ws_sum.cell(row=i, column=2).value = dtype
        ws_sum.cell(row=i, column=3).value = ntables
        ws_sum.cell(row=i, column=4).value = nrows
        _style_data_row(ws_sum, i, 4, alternate=(i % 2 == 0))

    _auto_fit_columns(ws_sum)

    wb.save(output_path)
    print(f"\n✅  Excel workbook saved to: {output_path}")
    print(f"   Processed {len(pdf_files)} PDF(s).")


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="Convert PDF quotes/invoices to a structured Excel workbook."
    )
    parser.add_argument(
        "pdf_directory",
        help="Path to the folder containing the PDF files.",
    )
    parser.add_argument(
        "output",
        nargs="?",
        default=None,
        help="Output Excel file path (default: <pdf_directory>/output.xlsx).",
    )
    args = parser.parse_args()

    output_path = args.output or str(Path(args.pdf_directory).resolve() / "output.xlsx")
    build_excel(args.pdf_directory, output_path)


if __name__ == "__main__":
    main()
